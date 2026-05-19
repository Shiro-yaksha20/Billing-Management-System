"""Report service for exports."""

from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..infrastructure.logging import logger
from ..repositories.bill_repository import BillRepository
from .settings_service import SettingsService


class ReportService:
    """Service for bill export operations."""

    def __init__(
        self,
        bill_repo: BillRepository,
        settings_service: SettingsService | None = None,
    ) -> None:
        self._bill_repo = bill_repo
        self._settings_service = settings_service

    def _currency_symbol(self) -> str:
        if not self._settings_service:
            return "\u20B9"
        return self._settings_service.get_setting("currency_symbol", "\u20B9") or "\u20B9"

    def export_bills(
        self,
        output_path: str,
        start_date: datetime | None = None,
        end_date: datetime | None = None,
        customer_id: int | None = None,
    ) -> bool:
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Bills Export"
            headers = [
                "Bill #",
                "Date",
                "Time",
                "Customer Name",
                "Customer Phone",
                "Staff",
                "Service",
                "Variant",
                "Qty",
                "Unit Price",
                "Line Total",
                "Subtotal",
                "Discount",
                "Tax %",
                "Tax Amount",
                "Total",
                "Payment Method",
                "Payment Status",
                "Transaction ID",
            ]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

            row = 2
            bills = self._bill_repo.find_for_export(
                start_date=start_date,
                end_date=end_date,
                customer_id=customer_id,
            )
            for bill in bills:
                for item in bill.items:
                    ws.cell(row=row, column=1, value=bill.bill_number or bill.id)
                    ws.cell(
                        row=row,
                        column=2,
                        value=bill.bill_datetime.strftime("%Y-%m-%d") if bill.bill_datetime else "",
                    )
                    ws.cell(
                        row=row,
                        column=3,
                        value=bill.bill_datetime.strftime("%H:%M") if bill.bill_datetime else "",
                    )
                    ws.cell(row=row, column=4, value=bill.customer.name if bill.customer else "")
                    ws.cell(row=row, column=5, value=bill.customer.phone if bill.customer else "")
                    ws.cell(row=row, column=6, value=bill.staff.name if bill.staff else "")
                    ws.cell(row=row, column=7, value=item.service.name if item.service else "")
                    ws.cell(row=row, column=8, value=item.service.variant if item.service else "")
                    ws.cell(row=row, column=9, value=item.quantity)
                    ws.cell(row=row, column=10, value=str(item.unit_price or Decimal("0")))
                    ws.cell(row=row, column=11, value=str(item.line_total or Decimal("0")))
                    ws.cell(row=row, column=12, value=str(bill.subtotal or Decimal("0")))
                    ws.cell(
                        row=row,
                        column=13,
                        value=str(bill.discount_amount or Decimal("0")),
                    )
                    ws.cell(row=row, column=14, value=str(bill.tax_percent or Decimal("0")))
                    ws.cell(row=row, column=15, value=str(bill.tax_amount or Decimal("0")))
                    ws.cell(row=row, column=16, value=str(bill.total or Decimal("0")))
                    ws.cell(row=row, column=17, value=bill.payment_method or "")
                    ws.cell(
                        row=row,
                        column=18,
                        value=getattr(bill, "payment_status", "Paid") or "Paid",
                    )
                    ws.cell(
                        row=row,
                        column=19,
                        value=getattr(bill, "transaction_id", "") or "",
                    )
                    row += 1
            wb.save(output_path)
            logger.info("Exported %s rows to %s", row - 2, output_path)
            return True
        except Exception as exc:
            logger.error("Excel export failed: %s", exc, exc_info=True)
            return False

    def export_bills_csv(
        self,
        output_path: str,
        start_date: datetime | None = None,
        end_date: datetime | None = None,
        customer_id: int | None = None,
    ) -> bool:
        """Export bills to CSV format."""
        try:
            headers = [
                "Bill #",
                "Date",
                "Time",
                "Customer Name",
                "Customer Phone",
                "Staff",
                "Service",
                "Variant",
                "Qty",
                "Unit Price",
                "Line Total",
                "Subtotal",
                "Discount",
                "Tax %",
                "Tax Amount",
                "Total",
                "Payment Method",
                "Payment Status",
                "Transaction ID",
            ]

            bills = self._bill_repo.find_for_export(
                start_date=start_date,
                end_date=end_date,
                customer_id=customer_id,
            )
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)

            with output_file.open("w", newline="", encoding="utf-8") as file_handle:
                writer = csv.writer(file_handle)
                writer.writerow(headers)
                for bill in bills:
                    for item in bill.items:
                        writer.writerow(
                            [
                                bill.bill_number or bill.id,
                                bill.bill_datetime.strftime("%Y-%m-%d") if bill.bill_datetime else "",
                                bill.bill_datetime.strftime("%H:%M") if bill.bill_datetime else "",
                                bill.customer.name if bill.customer else "",
                                bill.customer.phone if bill.customer else "",
                                bill.staff.name if bill.staff else "",
                                item.service.name if item.service else "",
                                item.service.variant if item.service else "",
                                item.quantity,
                                str(item.unit_price or Decimal("0")),
                                str(item.line_total or Decimal("0")),
                                str(bill.subtotal or Decimal("0")),
                                str(bill.discount_amount or Decimal("0")),
                                str(bill.tax_percent or Decimal("0")),
                                str(bill.tax_amount or Decimal("0")),
                                str(bill.total or Decimal("0")),
                                bill.payment_method or "",
                                getattr(bill, "payment_status", "Paid") or "Paid",
                                getattr(bill, "transaction_id", "") or "",
                            ]
                        )
            logger.info("Exported CSV to %s", output_path)
            return True
        except Exception as exc:
            logger.error("CSV export failed: %s", exc, exc_info=True)
            return False

    def export_bills_pdf_summary(
        self,
        output_path: str,
        start_date: datetime | None = None,
        end_date: datetime | None = None,
        customer_id: int | None = None,
    ) -> bool:
        """Export a PDF summary report for bills."""
        try:
            bills = list(
                self._bill_repo.find_for_export(
                    start_date=start_date,
                    end_date=end_date,
                    customer_id=customer_id,
                )
            )
            total_bills = len(bills)
            currency_symbol = self._currency_symbol()
            total_amount = sum((Decimal(bill.total or 0) for bill in bills), Decimal("0"))
            pending_amount = sum(
                (Decimal(bill.total or 0) for bill in bills if bill.payment_status == "Pending"),
                Decimal("0"),
            )

            doc = SimpleDocTemplate(output_path, pagesize=letter)
            styles = getSampleStyleSheet()
            story = [Paragraph("Bills Summary Report", styles["Heading1"])]

            date_range = "All Dates"
            if start_date or end_date:
                start = start_date.strftime("%Y-%m-%d") if start_date else "Start"
                end = end_date.strftime("%Y-%m-%d") if end_date else "End"
                date_range = f"{start} to {end}"
            story.append(Paragraph(f"Date Range: {date_range}", styles["Normal"]))
            story.append(Paragraph(f"Total Bills: {total_bills}", styles["Normal"]))
            story.append(
                Paragraph(
                    f"Total Amount: {currency_symbol}{total_amount:,.2f}",
                    styles["Normal"],
                )
            )
            story.append(
                Paragraph(
                    f"Pending Amount: {currency_symbol}{pending_amount:,.2f}",
                    styles["Normal"],
                )
            )
            story.append(Spacer(1, 12))

            max_bills = 50
            truncated = total_bills > max_bills
            bills = bills[:max_bills]
            if truncated:
                story.append(
                    Paragraph(
                        f"Note: Showing first {max_bills} of {total_bills} bills.",
                        styles["Normal"],
                    )
                )
                story.append(Spacer(1, 12))

            table_data = [["Bill #", "Date", "Customer", "Total", "Status"]]
            for bill in bills:
                table_data.append(
                    [
                        bill.bill_number or bill.id,
                        bill.bill_datetime.strftime("%Y-%m-%d") if bill.bill_datetime else "",
                        bill.customer.name if bill.customer else "",
                        f"{currency_symbol}{Decimal(bill.total or 0):,.2f}",
                        bill.payment_status or "",
                    ]
                )

            table = Table(table_data)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                        ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ]
                )
            )
            story.append(table)
            doc.build(story)
            logger.info("Exported PDF summary to %s", output_path)
            return True
        except Exception as exc:
            logger.error("PDF summary export failed: %s", exc, exc_info=True)
            return False
