import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from .models import Bill
from .database import db_session
from .utils import logger

def export_bills_to_excel(output_path: str, start_date: datetime = None, end_date: datetime = None, customer_id: int = None) -> bool:
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Bills Export"
        headers = [
            "Bill #", "Date", "Time", "Customer Name", "Customer Phone",
            "Staff", "Service", "Variant", "Qty", "Unit Price", "Line Total",
            "Subtotal", "Discount", "Tax %", "Tax Amount", "Total",
            "Payment Method", "Payment Status", "Transaction ID"
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        row = 2
        with db_session() as db:
            query = db.query(Bill)
            if start_date:
                query = query.filter(Bill.bill_datetime >= start_date)
            if end_date:
                query = query.filter(Bill.bill_datetime <= end_date)
            if customer_id:
                query = query.filter(Bill.customer_id == customer_id)
            bills = query.order_by(Bill.bill_datetime.desc()).all()
            for bill in bills:
                for item in bill.items:
                    ws.cell(row=row, column=1, value=bill.bill_number or bill.id)
                    ws.cell(row=row, column=2, value=bill.bill_datetime.strftime('%Y-%m-%d') if bill.bill_datetime else '')
                    ws.cell(row=row, column=3, value=bill.bill_datetime.strftime('%H:%M') if bill.bill_datetime else '')
                    ws.cell(row=row, column=4, value=bill.customer.name if bill.customer else '')
                    ws.cell(row=row, column=5, value=bill.customer.phone if bill.customer else '')
                    ws.cell(row=row, column=6, value=bill.staff.name if bill.staff else '')
                    ws.cell(row=row, column=7, value=item.service.name if item.service else '')
                    ws.cell(row=row, column=8, value=item.service.variant if item.service else '')
                    ws.cell(row=row, column=9, value=item.quantity)
                    ws.cell(row=row, column=10, value=float(item.unit_price) if item.unit_price else 0)
                    ws.cell(row=row, column=11, value=float(item.line_total) if item.line_total else 0)
                    ws.cell(row=row, column=12, value=float(bill.subtotal) if bill.subtotal else 0)
                    ws.cell(row=row, column=13, value=float(bill.discount_amount) if bill.discount_amount else 0)
                    ws.cell(row=row, column=14, value=float(bill.tax_percent) if bill.tax_percent else 0)
                    ws.cell(row=row, column=15, value=float(bill.tax_amount) if bill.tax_amount else 0)
                    ws.cell(row=row, column=16, value=float(bill.total) if bill.total else 0)
                    ws.cell(row=row, column=17, value=bill.payment_method or '')
                    ws.cell(row=row, column=18, value=getattr(bill, 'payment_status', 'Paid') or 'Paid')
                    ws.cell(row=row, column=19, value=getattr(bill, 'transaction_id', '') or '')
                    row += 1
        wb.save(output_path)
        logger.info(f"Exported {row-2} rows to {output_path}")
        return True
    except Exception as e:
        logger.error(f"Excel export failed: {e}")
        return False
